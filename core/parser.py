import re
from io import BytesIO
from typing import Dict, Any, List, Optional
import pandas as pd
from pypdf import PdfReader
from docx import Document

def extract_text_from_pdf(content: bytes) -> str:
    """Extract text from PDF file bytes."""
    try:
        reader = PdfReader(BytesIO(content))
        text_pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(text_pages).strip()
    except Exception as e:
        return f"[PDF Extraction Error: {str(e)}]"

def extract_text_from_docx(content: bytes) -> str:
    """Extract text from DOCX file bytes."""
    try:
        doc = Document(BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs).strip()
    except Exception as e:
        return f"[DOCX Extraction Error: {str(e)}]"

def extract_text_from_txt(content: bytes) -> str:
    """Extract text from TXT file bytes with fallback encodings."""
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="ignore")

def extract_from_excel(content: bytes) -> List[Dict[str, Any]]:
    """
    Extract structured candidate records from an Excel workbook (.xlsx).
    Normalizes column names matching original dataset schemas.
    """
    df = pd.read_excel(BytesIO(content))
    df.columns = df.columns.str.strip()
    df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
    
    candidates = []
    for _, row in df.iterrows():
        record = {str(k): (None if pd.isna(v) else v) for k, v in row.to_dict().items()}
        name = str(record.get('NAME', '') or record.get('name', ''))
        email = str(record.get('EMAILID', '') or record.get('email', '') or record.get('EMAIL', ''))
        edu = str(record.get('EDU', '') or record.get('edu', ''))
        yoe = str(record.get('YOE', '') if record.get('YOE') is not None else (record.get('yoe', '') or ''))
        tech = str(record.get('TECH KN.', '') or record.get('tech', ''))
        proj = str(record.get('HANDS ON PROJECT', '') or record.get('project', ''))
        cert = str(record.get('CERTIFICATIONS', '') or record.get('certifications', ''))
        
        parts = [
            f"Candidate: {name}" if name and name != 'None' else "",
            f"Email: {email}" if email and email != 'None' else "",
            f"Education: {edu}" if edu and edu != 'None' else "",
            f"Experience: {yoe} years of experience" if yoe and yoe != 'None' else "",
            f"Technical Skills: {tech}" if tech and tech != 'None' else "",
            f"Projects: {proj}" if proj and proj != 'None' else "",
            f"Certifications: {cert}" if cert and cert != 'None' else ""
        ]
        synthesized = "\n".join(p for p in parts if p).strip()
        existing_combined = str(record.get('combined_text', '') or '')
        if existing_combined and existing_combined != 'None' and existing_combined != synthesized:
            record['combined_text'] = synthesized + "\n" + existing_combined
        else:
            record['combined_text'] = synthesized or "Candidate Profile"
        candidates.append(record)
    return candidates

def parse_resume_content(content: bytes, filename: str) -> Dict[str, Any]:
    """Unified entry point for resume extraction across PDF, DOCX, TXT, and XLSX."""
    fn = filename.lower()
    
    if fn.endswith(".xlsx"):
        candidates = extract_from_excel(content)
        # For single resume flow, create a synthesized profile string from the first row
        first_text = candidates[0].get('combined_text', '') if candidates else ""
        return {
            "type": "excel",
            "candidates": candidates,
            "raw_text": first_text or f"Loaded Excel with {len(candidates)} candidates.",
            "filename": filename
        }
    elif fn.endswith(".pdf"):
        text = extract_text_from_pdf(content)
        return {"type": "pdf", "raw_text": text, "filename": filename}
    elif fn.endswith(".docx"):
        text = extract_text_from_docx(content)
        return {"type": "docx", "raw_text": text, "filename": filename}
    elif fn.endswith(".txt"):
        text = extract_text_from_txt(content)
        return {"type": "txt", "raw_text": text, "filename": filename}
    else:
        # Default fallback attempt as text
        text = extract_text_from_txt(content)
        return {"type": "unknown", "raw_text": text, "filename": filename}

def extract_candidate_metadata(text: str) -> Dict[str, Any]:
    """Lightweight deterministic regex extraction for email, phone, education and YOE hints."""
    email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
    phone_match = re.search(r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text)
    
    # Degree detection
    degree = "UG"
    lower_text = text.lower()
    if "phd" in lower_text or "p.hd" in lower_text or "doctorate" in lower_text:
        degree = "PhD"
    elif any(k in lower_text for k in ["master", "m.tech", "ms ", "m.s.", "pg ", "post graduate", "mba"]):
        degree = "PG"
    elif any(k in lower_text for k in ["bachelor", "b.tech", "bs ", "b.s.", "ug ", "undergraduate"]):
        degree = "UG"
        
    # YOE detection heuristic
    yoe = 0
    yoe_matches = re.findall(r'(\d+)\+?\s*(?:years?|yrs?)(?:\s+of)?\s+experience', lower_text)
    if yoe_matches:
        try:
            yoe = max([int(m) for m in yoe_matches])
        except ValueError:
            yoe = 0
            
    return {
        "email": email_match.group(0) if email_match else None,
        "phone": phone_match.group(0) if phone_match else None,
        "detected_degree": degree,
        "detected_yoe": yoe
    }
